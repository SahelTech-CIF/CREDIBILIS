"""
Lecteurs multi-sources du moteur de collecte CREDIBILIS.
Extrait les flux bruts sous forme d'objets `RawRecord`.
Ne dépend d'aucun framework web.
"""

from __future__ import annotations

import csv
import io
import json
import os
from typing import Any, Dict, Iterator, List, Optional, Tuple, Union

from .contracts import RawRecord
from .enums import SourceKind


class BaseReader:
    """Interface de base pour les lecteurs de fichiers."""
    def read_records(self, file_source: Union[str, bytes, io.IOBase], source_name: str, sheet_name: Optional[str] = None) -> Iterator[RawRecord]:
        raise NotImplementedError


class CSVReader(BaseReader):
    """Lecteur tolérant pour fichiers CSV."""

    def inspect(self, file_source: Union[str, bytes, io.IOBase], source_name: str) -> Dict[str, Any]:
        text = self._to_text(file_source)
        dialect, delimiter = self._detect_dialect(text[:4096])
        f = io.StringIO(text)
        reader = csv.reader(f, delimiter=delimiter)
        try:
            headers = [h.strip() for h in next(reader, []) if h.strip()]
        except Exception:
            headers = []
        rows_sample = []
        row_count = 0
        for r in reader:
            row_count += 1
            if len(rows_sample) < 5:
                rows_sample.append(r)
        return {
            "sheets": [{"name": "default", "columns": headers, "rows": row_count}],
            "delimiter": delimiter,
            "sample": rows_sample,
        }

    def read_records(self, file_source: Union[str, bytes, io.IOBase], source_name: str, sheet_name: Optional[str] = None) -> Iterator[RawRecord]:
        text = self._to_text(file_source)
        _, delimiter = self._detect_dialect(text[:4096])
        f = io.StringIO(text)
        reader = csv.DictReader(f, delimiter=delimiter)
        for idx, row in enumerate(reader, start=1):
            cleaned = {k.strip(): v.strip() if isinstance(v, str) else v for k, v in row.items() if k}
            yield RawRecord(row_number=idx, values=cleaned, source_name=source_name, sheet_name=sheet_name or "default")

    def _to_text(self, file_source: Union[str, bytes, io.IOBase]) -> str:
        if isinstance(file_source, str):
            if os.path.exists(file_source):
                with open(file_source, "rb") as f:
                    content = f.read()
            else:
                content = file_source.encode("utf-8")
        elif isinstance(file_source, bytes):
            content = file_source
        else:
            content = file_source.read()
            if isinstance(content, str):
                return content

        for encoding in ("utf-8-sig", "utf-8", "latin-1", "cp1252"):
            try:
                return content.decode(encoding)
            except UnicodeDecodeError:
                continue
        return content.decode("utf-8", errors="replace")

    def _detect_dialect(self, sample: str) -> Tuple[Any, str]:
        for delim in (";", ",", "\t", "|"):
            if delim in sample:
                return None, delim
        return None, ","


class ExcelReader(BaseReader):
    """Lecteur pour classeurs Microsoft Excel (.xlsx, .xlsm)."""

    def inspect(self, file_source: Union[str, bytes, io.IOBase], source_name: str) -> Dict[str, Any]:
        import openpyxl

        wb = self._load_workbook(file_source)
        sheets_info = []
        for name in wb.sheetnames:
            ws = wb[name]
            headers = []
            for cell in ws[1]:
                val = str(cell.value or "").strip()
                if val:
                    headers.append(val)
            sheets_info.append({
                "name": name,
                "columns": headers,
                "rows": max(0, ws.max_row - 1),
            })
        return {"sheets": sheets_info}

    def read_records(self, file_source: Union[str, bytes, io.IOBase], source_name: str, sheet_name: Optional[str] = None) -> Iterator[RawRecord]:
        wb = self._load_workbook(file_source)
        selected_sheet = sheet_name if sheet_name and sheet_name in wb.sheetnames else wb.sheetnames[0]
        ws = wb[selected_sheet]

        headers = [str(cell.value or "").strip() for cell in ws[1]]
        for row_idx in range(2, ws.max_row + 1):
            row_cells = [ws.cell(row=row_idx, column=col_idx).value for col_idx in range(1, len(headers) + 1)]
            if not any(v is not None and str(v).strip() != "" for v in row_cells):
                continue
            row_dict = {headers[i]: row_cells[i] for i in range(len(headers)) if headers[i]}
            yield RawRecord(row_number=row_idx, values=row_dict, source_name=source_name, sheet_name=selected_sheet)

    def _load_workbook(self, file_source: Union[str, bytes, io.IOBase]):
        import openpyxl

        if isinstance(file_source, bytes):
            return openpyxl.load_workbook(io.BytesIO(file_source), data_only=True, read_only=False)
        elif isinstance(file_source, str) and os.path.exists(file_source):
            return openpyxl.load_workbook(file_source, data_only=True, read_only=False)
        elif hasattr(file_source, "read"):
            data = file_source.read()
            if isinstance(data, str):
                data = data.encode("utf-8")
            return openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=False)
        else:
            raise ValueError("Source Excel invalide ou introuvable.")


class JSONReader(BaseReader):
    """Lecteur pour fichiers JSON (liste d'enregistrements)."""

    def inspect(self, file_source: Union[str, bytes, io.IOBase], source_name: str) -> Dict[str, Any]:
        data = self._load(file_source)
        if isinstance(data, list) and len(data) > 0 and isinstance(data[0], dict):
            headers = list(data[0].keys())
            return {"sheets": [{"name": "records", "columns": headers, "rows": len(data)}]}
        return {"sheets": [{"name": "records", "columns": [], "rows": 0}]}

    def read_records(self, file_source: Union[str, bytes, io.IOBase], source_name: str, sheet_name: Optional[str] = None) -> Iterator[RawRecord]:
        data = self._load(file_source)
        if isinstance(data, dict):
            data = [data]
        for idx, item in enumerate(data, start=1):
            if isinstance(item, dict):
                yield RawRecord(row_number=idx, values=item, source_name=source_name, sheet_name="records")

    def _load(self, file_source: Union[str, bytes, io.IOBase]) -> Any:
        if isinstance(file_source, str):
            if os.path.exists(file_source):
                with open(file_source, "r", encoding="utf-8") as f:
                    return json.load(f)
            return json.loads(file_source)
        elif isinstance(file_source, bytes):
            return json.loads(file_source.decode("utf-8"))
        else:
            return json.load(file_source)


def reader_for(source_name: str, format_hint: Optional[SourceKind] = None) -> BaseReader:
    """Fabrique de lecteurs selon le format détecté."""
    name_lower = source_name.lower()
    if format_hint == SourceKind.XLSX or name_lower.endswith((".xlsx", ".xlsm")):
        return ExcelReader()
    if format_hint == SourceKind.JSON or name_lower.endswith(".json"):
        return JSONReader()
    return CSVReader()
